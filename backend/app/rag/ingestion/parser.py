import csv
import io
import zipfile
from pathlib import Path
from app.rag.errors import RAGError
from langchain_core.documents import Document as LCDocument

SUPPORTED = {".pdf", ".txt", ".md", ".markdown", ".docx", ".csv"}

def parse_document(data: bytes, extension: str) -> list[LCDocument]:
    extension = extension.lower()
    if extension not in SUPPORTED:
        raise RAGError("Unsupported document type", 415)
    if not data:
        raise RAGError("Document is empty", 422)
    try:
        if extension == ".pdf":
            import pymupdf
            if not data.startswith(b"%PDF-"):
                raise ValueError("Invalid PDF signature")
            with pymupdf.open(stream=data, filetype="pdf") as pdf:
                if pdf.needs_pass:
                    raise RAGError("Password-protected PDF is unsupported", 422)
                sections = [LCDocument(page_content=page.get_text(), metadata={"page_number": i+1}) for i, page in enumerate(pdf)]
            if not any(s.page_content.strip() for s in sections):
                raise RAGError("PDF has no extractable text; OCR is required", 422)
        elif extension == ".docx":
            from docx import Document
            # Reject compressed bombs before expanding the archive.
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise RAGError("DOCX expanded content is too large", 413)
            doc = Document(io.BytesIO(data))
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            parts = []
            for block in doc.iter_inner_content():
                if isinstance(block, Paragraph):
                    parts.append(block.text)
                elif isinstance(block, Table):
                    parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in block.rows))
            sections = [LCDocument(page_content="\n\n".join(parts))]
        else:
            text = data.decode("utf-8-sig")
            if "\x00" in text or any(ord(c) < 32 and c not in "\n\r\t" for c in text):
                raise ValueError("Binary content is not supported")
            if extension == ".csv":
                rows = csv.reader(io.StringIO(text), strict=True)
                headers = next(rows, [])
                if not headers or not any(h.strip() for h in headers):
                    raise ValueError("CSV headers are required")
                sections = []
                for number, row in enumerate(rows, 2):
                    if not any(cell.strip() for cell in row):
                        continue
                    if len(row) != len(headers):
                        raise ValueError("CSV row length does not match headers")
                    sections.append(LCDocument(page_content=" | ".join(f"{h}: {v}" for h, v in zip(headers, row)), metadata={"row_number": number, "headers": headers}))
            else:
                sections = [LCDocument(page_content=text)]
        if not any(s.page_content.strip() for s in sections):
            raise RAGError("Document has no usable text", 422)
        return sections
    except RAGError:
        raise
    except Exception as exc:
        raise RAGError("Invalid document content; use an unencrypted supported file (UTF-8 for text/CSV)", 422) from exc
