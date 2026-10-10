import asyncio
import io
import json
from concurrent.futures import ThreadPoolExecutor
import numpy as np
import pymupdf
import pytest
from langchain_core.messages import HumanMessage
from docx import Document as Docx
from pydantic import ValidationError
from app.rag.config import RAGSettings
from app.rag.errors import RAGError
from app.rag.schemas import QueryRequest
from app.rag.ingestion.chunker import split_text
from app.rag.ingestion.parser import parse_document
from app.rag.generation.prompts import build_context, validate_citations
from app.rag.generation.llm_provider import OpenAICompatibleProvider
from app.rag.pipeline.rag_pipeline import RAGPipeline
from app.rag.vectorstore.faiss_store import FaissStore
from app.rag.embeddings.sentence_transformer import SentenceTransformerEmbeddings
from conftest import FakeEmbeddings

@pytest.mark.parametrize("strategy", ["fixed", "paragraph"])
@pytest.mark.parametrize("text", ["", "   ", "abc", "a"*301, "αβ🌍"*90, "one paragraph\n\n"+"long "*90+"\n\nlast"])
def test_chunking(text, strategy):
    chunks = split_text(text, 30, 7, strategy)
    assert all(c.strip() and len(c) <= 30 for c in chunks)
    if not text.strip():
        assert chunks == []
        return
    # Reconstruct coverage using source ranges, including overlapping whitespace.
    cursor = 0
    covered = set()
    for chunk in chunks:
        position = text.find(chunk, cursor)
        assert position >= 0
        covered.update(range(position, position+len(chunk)))
        cursor = max(position+len(chunk)-7, position+1)
    assert all(i in covered for i,c in enumerate(text) if not c.isspace())

@pytest.mark.parametrize("size,overlap", [(0,0),(10,10),(10,-1)])
def test_invalid_chunking(size, overlap):
    with pytest.raises(ValueError):
        split_text("abc", size, overlap)

def test_parsing_text_and_csv():
    assert parse_document("héllo\n\nworld".encode(), ".txt")[0].page_content == "héllo\n\nworld"
    assert parse_document(b"# Title\n\nBody", ".md")[0].page_content.startswith("#")
    csv = parse_document(b'name,value\n"pump, blue",42\n', '.csv')
    assert csv[0].metadata == dict(row_number=2, headers=['name', 'value'])
    assert 'name: pump, blue' in csv[0].page_content

@pytest.mark.parametrize("data,extension", [(b"",'.txt'),(b"\x00",'.txt'),(b'abc','.pdf'),(b'abc','.exe'),(b'a,b\n1\n','.csv'),(b'a,b\n','.csv')])
def test_invalid_documents(data, extension):
    with pytest.raises(RAGError):
        parse_document(data, extension)

def test_pdf_pages_and_ocr():
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50,50), "pump first page")
        pdf.new_page().insert_text((50,50), "motor second page")
        sections = parse_document(pdf.tobytes(), '.pdf')
    assert [s.metadata['page_number'] for s in sections] == [1,2]
    with pymupdf.open() as pdf:
        pdf.new_page()
        with pytest.raises(RAGError, match='OCR'):
            parse_document(pdf.tobytes(), '.pdf')

def test_docx_tables():
    doc = Docx()
    doc.add_paragraph('pump manual')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0,0).text = 'motor'
    table.cell(0,1).text = '42'
    stream = io.BytesIO()
    doc.save(stream)
    text = parse_document(stream.getvalue(), '.docx')[0].page_content
    assert 'pump manual' in text and 'motor | 42' in text

def test_storage_retrieval_reload_delete(pipeline):
    first = pipeline.loader.ingest(b'pump pump maintenance', '../../pump.txt')
    second = pipeline.loader.ingest(b'motor motor motor', 'motor.txt')
    calls = pipeline.embeddings.calls
    duplicate = pipeline.loader.ingest(b'pump pump maintenance', 'same.txt')
    assert duplicate.deduplicated and pipeline.embeddings.calls == calls
    query = QueryRequest(query='pump', top_k=2)
    hits = pipeline.retriever.retrieve(query)
    assert hits[0].document_id == first.document.document_id
    assert hits[0].filename == 'pump.txt' and hits[0].content_hash == first.document.content_hash
    assert hits[0].score >= hits[1].score
    restricted = pipeline.retriever.retrieve(QueryRequest(query='pump', document_ids=[second.document.document_id]))
    assert len(restricted) == 1 and restricted[0].document_id == second.document.document_id
    assert pipeline.retriever.retrieve(QueryRequest(query='pump', document_ids=[])) == []
    assert len(pipeline.retriever.retrieve(QueryRequest(query='pump', similarity_threshold=.9))) == 1
    reloaded = RAGPipeline(pipeline.settings, embeddings=FakeEmbeddings())
    assert reloaded.retriever.retrieve(query)[0].chunk_id == hits[0].chunk_id
    reloaded.loader.delete(first.document.document_id)
    assert len(pipeline.retriever.retrieve(query)) == 1
    pipeline.loader.delete(second.document.document_id)
    assert reloaded.retriever.retrieve(query) == []
    with pytest.raises(RAGError, match='not found'):
        pipeline.loader.delete(second.document.document_id)

def test_empty_and_dimension_validation(pipeline):
    assert pipeline.retriever.retrieve(QueryRequest(query='pump')) == []
    assert pipeline.embeddings.calls == 0
    pipeline.loader.ingest(b'pump', 'p.txt')
    with pipeline.store.transaction():
        with pytest.raises(RAGError, match='dimension'):
            pipeline.store.search(np.ones((1,2)), 1)
    with pytest.raises(RAGError):
        FaissStore.normalize(np.zeros((1,3)))

def test_metadata_pdf(pipeline):
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50,50), 'pump')
        pipeline.loader.ingest(pdf.tobytes(), 'manual.pdf')
    hit = pipeline.retriever.retrieve(QueryRequest(query='pump'))[0]
    assert hit.page_number == 1 and hit.chunk_index == 0 and hit.document_type == 'pdf'

def test_index_incompatibility_and_corruption(pipeline):
    pipeline.loader.ingest(b'pump', 'p.txt')
    class OtherEmbeddings(FakeEmbeddings):
        identity = 'other'
    other = FaissStore(pipeline.settings.rag_index_dir, OtherEmbeddings(), pipeline.store.chunk_config)
    with pytest.raises(RAGError, match='incompatible'):
        with other.transaction():
            pass
    folder = pipeline.store.directory / (pipeline.store.directory/'CURRENT').read_text()
    (folder/'vectors.faiss').write_bytes(b'corrupted')
    with pytest.raises(RAGError, match='checksum'):
        with pipeline.store.transaction():
            pass

def test_citations_context_and_generation(pipeline):
    pipeline.loader.ingest(b'pump manual', 'p.txt')
    hits = pipeline.retriever.retrieve(QueryRequest(query='pump'))
    context, included = build_context(hits, 1000)
    assert len(context) <= 1000 and included == hits
    identifier = hits[0].chunk_id
    assert validate_citations(f'Pump [{identifier}]', hits)[0].chunk_id == identifier
    for answer in ('claim [chunk_invalid]', 'uncited claim'):
        with pytest.raises(RAGError):
            validate_citations(answer, hits)
    class MockLLM:
        async def generate(self, messages):
            assert identifier in messages[1].content
            return f'Pump manual [{identifier}]'
    pipeline.llm = MockLLM()
    response = asyncio.run(pipeline.query(QueryRequest(query='pump')))
    assert response.citations[0].document_id == hits[0].document_id
    assert response.metadata['latency_ms'] >= 0

def test_missing_llm_and_settings(pipeline):
    with pytest.raises(RAGError, match='LLM_BASE_URL'):
        asyncio.run(pipeline.query(QueryRequest(query='pump')))
    with pytest.raises(ValidationError):
        RAGSettings(_env_file=None, rag_chunk_size=10, rag_chunk_overlap=10)
    with pytest.raises(ValidationError):
        QueryRequest(query='  ')

def test_embedding_lazy_batch(monkeypatch):
    import sys
    import types
    loads = []
    class Model:
        def __init__(self, *args, **kwargs):
            loads.append(args)
        def encode(self, texts, **kwargs):
            assert kwargs['normalize_embeddings'] and kwargs['batch_size'] == 2
            return np.ones((len(texts),3))
    monkeypatch.setitem(sys.modules, 'sentence_transformers', types.SimpleNamespace(SentenceTransformer=Model))
    embedder = SentenceTransformerEmbeddings('test', 2)
    assert embedder.embed_documents([]) == [] and not loads
    assert np.asarray(embedder.embed_documents(['a','b'])).shape == (2,3)
    assert len(embedder.embed_query('query')) == 3
    assert len(loads) == 1

def test_api_and_existing_routes(client, pipeline):
    health = client.get('/health')
    assert health.status_code == 200 and 'ram' in health.json()
    assert 'x-inference-time-ms' in health.headers
    result = client.post('/api/v1/predict', json={'data':[1,2,3], 'metadata':{'test':True}})
    assert result.status_code == 200 and result.json()['input_mean'] == 2
    assert result.json()['metadata'] == {'test':True}
    cors = client.options('/health', headers={'Origin':'http://example.com','Access-Control-Request-Method':'GET'})
    assert cors.headers['access-control-allow-origin'] == '*'
    assert client.get('/api/rag/health').status_code == 200
    assert pipeline.embeddings.calls == 0
    assert client.post('/api/rag/retrieve', json={'query':' '}).status_code == 422
    assert client.post('/api/rag/retrieve', json={'query':'pump','top_k':0}).status_code == 422
    assert client.post('/api/rag/query', json={'query':'pump'}).status_code == 503
    assert client.post('/api/rag/documents/upload', files={'file':('a.exe',b'abc')}).status_code == 415
    assert client.post('/api/rag/documents/upload', files={'file':('a.txt',b'')}).status_code == 422
    upload = client.post('/api/rag/documents/upload', files={'file':('../../manual.txt',b'pump manual')})
    assert upload.status_code == 201
    identifier = upload.json()['document']['document_id']
    assert len(client.get('/api/rag/documents').json()) == 1
    assert client.get('/api/rag/documents/'+identifier).status_code == 200
    assert client.get('/api/rag/documents/missing').status_code == 404
    assert client.post('/api/rag/retrieve', json={'query':'pump'}).json()['chunks'][0]['document_id'] == identifier
    assert client.delete('/api/rag/documents/'+identifier).status_code == 204
    assert client.post('/api/rag/retrieve', json={'query':'pump'}).json() == {'chunks':[]}
    pipeline.settings.rag_enabled = False
    assert client.get('/health').status_code == 200
    assert client.get('/api/rag/health').json()['enabled'] is False
    assert client.get('/api/rag/documents').status_code == 503

def test_upload_limit(client, pipeline):
    pipeline.settings.rag_max_upload_mb = 1
    assert client.post('/api/rag/documents/upload', files={'file':('a.txt',b'a'*(1024*1024+1))}).status_code == 413

def test_concurrent_writers(pipeline):
    def ingest(index):
        independent = RAGPipeline(pipeline.settings, embeddings=FakeEmbeddings())
        return independent.loader.ingest(f'pump {index}'.encode(), f'{index}.txt')
    with ThreadPoolExecutor(max_workers=3) as executor:
        results = list(executor.map(ingest, range(6)))
    with pipeline.store.transaction():
        assert len(pipeline.store.documents) == len(results) == 6
        assert pipeline.store.index.ntotal == len(pipeline.store.chunks) == 6

def test_llm_http_mock(pipeline):
    import httpx
    calls = []
    def handler(request):
        calls.append(request)
        assert request.url.path == '/v1/chat/completions'
        assert json.loads(request.content)['model'] == 'test-model'
        if len(calls) == 1:
            return httpx.Response(429, headers={'retry-after-ms': '1'})
        return httpx.Response(200, json={'id':'x','object':'chat.completion','created':0,'model':'test-model','choices':[{'index':0,'finish_reason':'stop','message':{'role':'assistant','content':'evidence [chunk_abc]'}}]})
    pipeline.settings.llm_base_url = 'https://example.test/v1'
    pipeline.settings.llm_model = 'test-model'
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await OpenAICompatibleProvider(pipeline.settings, http_async_client=http).generate([HumanMessage('hello')])
    assert asyncio.run(run()) == 'evidence [chunk_abc]' and len(calls) == 2

def test_llm_errors_map_to_http_status(pipeline):
    import httpx
    pipeline.settings.llm_base_url = 'https://example.test/v1'
    pipeline.settings.llm_model = 'test-model'
    pipeline.settings.llm_max_retries = 0
    async def run(handler):
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await OpenAICompatibleProvider(pipeline.settings, http_async_client=http).generate([HumanMessage('hello')])
    def unreachable(request):
        raise httpx.ConnectError('down')
    for handler, status in ((lambda r: httpx.Response(500), 502), (lambda r: httpx.Response(400), 502), (unreachable, 504)):
        with pytest.raises(RAGError) as error:
            asyncio.run(run(handler))
        assert error.value.status_code == status

def test_langgraph_query_with_chat_model(pipeline):
    from langchain_core.language_models.fake_chat_models import FakeListChatModel
    pipeline.loader.ingest(b'pump manual', 'p.txt')
    identifier = pipeline.retriever.retrieve(QueryRequest(query='pump'))[0].chunk_id
    graph_pipeline = RAGPipeline(pipeline.settings, embeddings=FakeEmbeddings(), llm=FakeListChatModel(responses=[f'Pump [{identifier}]']))
    response = asyncio.run(graph_pipeline.query(QueryRequest(query='pump')))
    assert response.citations[0].chunk_id == identifier
    assert set(graph_pipeline.graph.get_graph().nodes) >= {'retrieve', 'rerank', 'build_context', 'generate', 'insufficient_evidence'}
    empty = asyncio.run(graph_pipeline.query(QueryRequest(query='pump', document_ids=[])))
    assert empty.citations == [] and 'insufficient' in empty.answer

def test_langchain_retriever_and_vectorstore(pipeline):
    pipeline.loader.ingest(b'pump pump maintenance', 'pump.txt')
    pipeline.loader.ingest(b'motor motor motor', 'motor.txt')
    documents = pipeline.retriever.invoke('pump', top_k=1)
    assert len(documents) == 1 and documents[0].page_content == 'pump pump maintenance'
    assert documents[0].metadata['filename'] == 'pump.txt' and documents[0].id.startswith('chunk_')
    scored = pipeline.store.similarity_search_with_score('motor', k=2)
    assert scored[0][0].metadata['filename'] == 'motor.txt' and scored[0][1] >= scored[1][1]
